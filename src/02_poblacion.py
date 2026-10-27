"""
02_poblacion.py
Attach DPA canton codes to the INEC population projections and validate
canton codes across the three sources (population, boundaries, homicides).

Input:  data/external/Tabulado_cantonal_2010-2035_rev2024.xlsx  (INEC, rev. 2024)
        data/external/cantones/cantones.shp                     (INEC boundaries)
        data/interim/homicidios_limpio.csv                      (from 01_limpieza.py)
Output: data/processed/poblacion_canton_anio.csv
        (canton x year, 2014-2026; aggregate data, safe to publish)

Run from anywhere:  python src/02_poblacion.py
"""

import sys
import unicodedata
from pathlib import Path

import geopandas as gpd
import openpyxl
import pandas as pd

# ---------------------------------------------------------------------------
# Paths and parameters
# ---------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[1]
POP_FILE = ROOT / "data" / "external" / "Tabulado_cantonal_2010-2035_rev2024.xlsx"
SHP_FILE = ROOT / "data" / "external" / "cantones" / "cantones.shp"
HOM_FILE = ROOT / "data" / "interim" / "homicidios_limpio.csv"
OUT_FILE = ROOT / "data" / "processed" / "poblacion_canton_anio.csv"

YEARS = list(range(2014, 2027))  # study period (2026 = partial year)

# One sheet per province, in DPA order. Province code = DPA code.
SHEET_TO_PROVINCE = {
    "Azuay": "01", "Bolivar": "02", "Cañar": "03", "Carchi": "04",
    "Cotopaxi": "05", "Chimborazo": "06", "El Oro": "07", "Esmeraldas": "08",
    "Guayas": "09", "Imbabura": "10", "Loja": "11", "Los Rios": "12",
    "Manabí": "13", "Morona": "14", "Napo": "15", "Pastaza": "16",
    "Pichincha": "17", "Tungurahua": "18", "Zamora": "19", "Galapagos": "20",
    "Sucumbios": "21", "Orellana": "22", "Santo Domingo": "23", "Santa Elena": "24",
}

# Population-table name -> boundary-layer name, when they differ (normalized).
# Fill only if the name check below reports unmatched cantons.
NAME_ALIASES = {}


def normalize(x) -> str:
    """Upper case, strip, remove accents (Ñ -> N)."""
    x = unicodedata.normalize("NFKD", str(x)).encode("ascii", "ignore").decode()
    return " ".join(x.upper().split())


def check(condition: bool, message: str) -> None:
    """Stop the pipeline if a verification fails."""
    if not condition:
        sys.exit(f"[FAIL] {message}")
    print(f"[ok]   {message}")


# ---------------------------------------------------------------------------
# 1. Read the population table (one sheet per province)
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(POP_FILE, read_only=True, data_only=True)
check(set(SHEET_TO_PROVINCE) == set(wb.sheetnames[1:]), "all 24 province sheets present")

records, province_totals = [], []
for sheet, prov_code in SHEET_TO_PROVINCE.items():
    rows = [[c for c in r if c is not None] for r in wb[sheet].iter_rows(values_only=True)]
    header = next(r for r in rows if len(r) == 26 and r[0] == 2010)  # 2010 ... 2035
    data = [r for r in rows if len(r) == 27 and isinstance(r[0], str)]
    # First data row = province total; the rest = cantons
    province_totals.append(dict(zip(["name"] + header, data[0])) | {"codigo_provincia": prov_code})
    for r in data[1:]:
        rec = dict(zip(["canton_inec"] + header, r))
        rec["codigo_provincia"] = prov_code
        records.append(rec)

pop = pd.DataFrame(records)
prov_tot = pd.DataFrame(province_totals)
check(len(pop) == 221, f"221 cantons in the population table ({len(pop)})")

# Internal consistency: cantons add up to their province total, every year
canton_sum = pop.groupby("codigo_provincia")[YEARS].sum()
check((canton_sum.values == prov_tot.set_index("codigo_provincia").loc[canton_sum.index, YEARS].values).all(),
      "cantons add up to the province totals (2014-2026)")

# ---------------------------------------------------------------------------
# 2. Attach DPA codes from the boundary layer (province code + canton name)
# ---------------------------------------------------------------------------
shp = gpd.read_file(SHP_FILE, encoding="utf-8")[["code", "name_es"]]
check(shp["code"].is_unique and len(shp) == 221, "221 unique codes in the boundary layer")
shp["codigo_provincia"] = shp["code"].str[:2]
shp["key"] = shp["codigo_provincia"] + "|" + shp["name_es"].map(normalize)

pop_name = pop["canton_inec"].map(normalize).replace(NAME_ALIASES)
pop["key"] = pop["codigo_provincia"] + "|" + pop_name
check(pop["key"].is_unique, "canton names are unique within each province")

pop = pop.merge(shp[["key", "code", "name_es"]], on="key", how="left")
unmatched = pop.loc[pop["code"].isna(), ["codigo_provincia", "canton_inec"]]
check(unmatched.empty, "every population canton matches a boundary polygon"
      + ("" if unmatched.empty else f"\n{unmatched.to_string()}"))
check(pop["code"].is_unique, "the match is one-to-one")

# ---------------------------------------------------------------------------
# 3. Validate homicide canton codes against the other two sources
# ---------------------------------------------------------------------------
hom = pd.read_csv(HOM_FILE, dtype={"codigo_canton": str})
hom_codes = set(hom["codigo_canton"])
check(hom_codes <= set(pop["code"]), f"all {len(hom_codes)} homicide canton codes have population")
no_cases = shp.loc[~shp["code"].isin(hom_codes), ["code", "name_es"]]
print(f"[info] cantons with no recorded homicides (true zeros): {len(no_cases)}")
print(no_cases.to_string(index=False))

# ---------------------------------------------------------------------------
# 4. Long format, study period only
# ---------------------------------------------------------------------------
out = (
    pop.rename(columns={"code": "codigo_canton", "name_es": "canton"})
    .melt(id_vars=["codigo_canton", "canton", "codigo_provincia"],
          value_vars=YEARS, var_name="anio", value_name="poblacion")
    .sort_values(["codigo_canton", "anio"])
    .reset_index(drop=True)
)
out["anio"] = out["anio"].astype(int)
out["poblacion"] = out["poblacion"].round().astype(int)

check(len(out) == 221 * len(YEARS), f"221 cantons x {len(YEARS)} years ({len(out):,} rows)")
check((out["poblacion"] > 0).all(), "every population value is positive")

national = out.groupby("anio")["poblacion"].sum()
print("\n[info] national population (sum of cantons):")
print(national.to_string())

# ---------------------------------------------------------------------------
# 5. Save
# ---------------------------------------------------------------------------
OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
out.to_csv(OUT_FILE, index=False, encoding="utf-8")
print(f"\nSaved {OUT_FILE.relative_to(ROOT)} ({len(out):,} rows)")